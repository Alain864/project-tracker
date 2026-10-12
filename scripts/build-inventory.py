"""Turn tracker-projects-v2.xlsx into src/inventory.json.

Each sheet is a project. Columns B–E are DEV, TEST, PREPROD, and PROD.
Re-run after the four empty projects are filled:

  python3 scripts/build-inventory.py /path/to/tracker-projects-v2.xlsx
"""

import csv
import io
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path("/Users/alaincl/Downloads/tracker-projects-v2.xlsx")
OUT = ROOT / "src" / "inventory.json"

SHEET_IDS = {
    "cloud-platform-cost-investigato": "cloud-platform-cost-investigator",
}
COMPONENTS = [
    "VM",
    "Pg",
    "Pg Vector",
    "Redis",
    "Langfuse",
    "Openshift",
    "Argocd",
    "Github-repo",
    "Github-teams",
    "Nesgen-secrets",
    "Keyvault",
    "Spns",
    "Spn-graph",
    "fe-spn",
    "fe-spn-graph",
    "DBX",
    "Security group /AAD group",
    "Persona groups",
    "Keyvault check role assigment",
]
COMPONENT_IDS = [
    "vm",
    "postgres",
    "pgvector",
    "redis",
    "langfuse",
    "openshift",
    "argocd",
    "github-repo",
    "github-teams",
    "nesgen",
    "keyvault",
    "spns",
    "spns-graph",
    "frontend-spns",
    "fe-spn-graph",
    "databricks",
    "aad-groups",
    "persona-groups",
    "keyvault-roles",
]
ENVS = ["DEV", "TEST", "PREPROD", "PROD"]
PROJECT_IDS = [
    "36411-aura-brandiq",
    "36519-aura-trike",
    "32938-aura-order-intel",
    "32936-aura-reliability",
    "33155-photon",
    "cloud-platform-knowledge",
    "cloud-platform-analyzer",
    "cloud-platform-cost-investigator",
]
PERM = re.compile(
    r"^(email|offline_access|openid|profile|User\.Read)\s+(Granted|Pending)(?:\s+for\s+(\S+))?$",
    re.I,
)
TEAM_MEMBER = re.compile(
    r"^(.+?)\s+[–—-]\s+(\d+)\s+members\s+[–—-]\s+Role:\s*(\w+)\s*$",
    re.I,
)


def clean(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\u00a0", " ").replace("\r\n", "\n").strip()


def lines_of(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def title_state(value: str) -> str:
    return "Pending" if value.lower() == "pending" else "Granted"


def parse_teams(text: str):
    rows_src = lines_of(text)
    if not rows_src:
        return None
    if rows_src[0].lower().startswith("team name"):
        rows = []
        for line in rows_src[1:]:
            parts = [part.strip() for part in re.split(r"\t+", line) if part.strip()]
            if len(parts) >= 2:
                rows.append({"name": parts[0], "role": parts[1]})
        return {"kind": "teams", "rows": rows} if rows else None
    parsed = []
    for line in rows_src:
        match = TEAM_MEMBER.match(line)
        if not match:
            return None
        parsed.append(
            {"name": match.group(1).strip(), "role": match.group(3), "members": match.group(2)}
        )
    return {"kind": "teams", "rows": parsed} if parsed else None


def parse_permissions(text: str):
    principal = None
    rows = []
    for line in lines_of(text):
        parts = [part.strip() for part in line.split("\t")]
        candidate = None
        if len(parts) >= 2 and parts[0] and parts[1]:
            match = re.match(r"(Granted|Pending)(?:\s+for\s+(\S+))?$", parts[1], re.I)
            if match and not parts[0].lower().startswith("microsoft graph"):
                candidate = (parts[0], title_state(match.group(1)), match.group(2))
        if candidate is None:
            match = PERM.match(re.sub(r"\s+", " ", line))
            if match:
                candidate = (match.group(1), title_state(match.group(2)), match.group(3))
        if candidate:
            scope, state, tenant = candidate
            row = {"scope": scope, "state": state}
            if tenant:
                row["tenant"] = tenant
            rows.append(row)
            continue
        if principal is None:
            principal = re.sub(r"\s+", " ", line).strip()
    if not rows:
        return None
    detail = {"kind": "permissions", "rows": rows}
    if principal:
        detail["principal"] = principal
    return detail


def role_rows_from(records: list[list[str]]):
    rows = []
    section = ""
    for record in records:
        cells = [(record[i].strip() if i < len(record) else "") for i in range(3)]
        name, kind, role = cells
        if not name or name.lower() in {"section / name", "name / id"}:
            continue
        if not kind and not role:
            section = name
            continue
        name_lines = [line.strip() for line in name.splitlines() if line.strip()]
        row = {"section": section, "name": name_lines[0], "type": kind, "role": role}
        if len(name_lines) > 1:
            row["id"] = name_lines[1]
        rows.append(row)
    return {"kind": "roles", "rows": rows} if rows else None


def parse_roles(text: str):
    first = lines_of(text)[:1]
    header = first[0].lower() if first else ""
    if "name / id" not in header and "section / name" not in header and "key vault" not in text.lower():
        return None
    if "," in text.splitlines()[0]:
        return role_rows_from(list(csv.reader(io.StringIO(text))))
    records = []
    for line in text.splitlines():
        if not line.strip():
            continue
        parts = line.split("\t")
        records.append(parts[:3])
    return role_rows_from(records)


def parse_list(text: str):
    items = []
    for line in lines_of(text):
        deprecated = bool(re.search(r"\(DEPRECATE\)\s*$", line, re.I))
        label = re.sub(r"\s*\(DEPRECATE\)\s*$", "", line, flags=re.I).strip()
        if label:
            item = {"label": label}
            if deprecated:
                item["deprecated"] = True
            items.append(item)
    if len(items) < 2:
        return None
    return {"kind": "list", "items": items}


def summarize(status: str, detail) -> str:
    if status == "N/A":
        return "Not required"
    if status == "Blocked":
        return "Missing"
    if status == "Not Started":
        return "Not recorded"
    if not detail:
        return "Present"
    if detail["kind"] == "text":
        return detail["value"]
    if detail["kind"] == "list":
        count = len(detail["items"])
        return f"{count} group" if count == 1 else f"{count} groups"
    if detail["kind"] == "teams":
        count = len(detail["rows"])
        return f"{count} team" if count == 1 else f"{count} teams"
    if detail["kind"] == "permissions":
        pending = [row for row in detail["rows"] if row["state"] == "Pending"]
        if pending:
            return f"{len(pending)} pending"
        count = len(detail["rows"])
        return f"{count} granted"
    count = len(detail["rows"])
    return f"{count} assignment" if count == 1 else f"{count} assignments"


def other_project(principal: str, project_id: str) -> str | None:
    for other in PROJECT_IDS:
        if other != project_id and other in principal:
            return other
    return None


def cell_from(text: str, recorded: bool, project_id: str):
    if not text:
        summary = "Awaiting data" if not recorded else "Not recorded"
        return {"status": "Not Started", "summary": summary, "detail": None, "blocker": None}
    folded = text.lower()
    if folded == "na":
        return {"status": "N/A", "summary": "Not required", "detail": None, "blocker": None}
    if folded == "missing":
        return {
            "status": "Blocked",
            "summary": "Missing",
            "detail": None,
            "blocker": {
                "severity": "High",
                "description": "Marked missing in the inventory.",
            },
        }
    if text in {"✓", "✔", "√"}:
        detail = {"kind": "text", "value": "Present"}
        return {"status": "Ready", "summary": "Present", "detail": detail, "blocker": None}

    detail = parse_teams(text) or parse_roles(text) or parse_permissions(text) or parse_list(text)
    if detail is None:
        detail = {"kind": "text", "value": re.sub(r"\s+", " ", text).strip()}
    pending = []
    if detail["kind"] == "permissions":
        pending = [row["scope"] for row in detail["rows"] if row["state"] == "Pending"]
    status = "In Progress" if pending else "Ready"
    blocker = None
    summary = summarize(status, detail)
    if pending:
        blocker = {
            "severity": "Medium",
            "description": "Pending consent: " + ", ".join(pending) + ".",
        }
    if detail and detail.get("kind") == "permissions":
        other = other_project(detail.get("principal") or "", project_id)
        if other:
            status = "Blocked"
            summary = other
            blocker = {
                "severity": "High",
                "description": (
                    f"These grants belong to {detail['principal']} ({other}), not this project."
                ),
            }
    return {
        "status": status,
        "summary": summary,
        "detail": detail,
        "blocker": blocker,
    }


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    workbook = openpyxl.load_workbook(source, data_only=False)
    projects = []
    for sheet in workbook.sheetnames:
        ws = workbook[sheet]
        project_id = SHEET_IDS.get(sheet, sheet)
        found = {
            clean(ws.cell(row, 1).value): row
            for row in range(1, ws.max_row + 1)
            if clean(ws.cell(row, 1).value)
        }
        raw = {}
        any_value = False
        for label, component_id in zip(COMPONENTS, COMPONENT_IDS):
            row = found.get(label)
            raw[component_id] = {}
            for index, env in enumerate(ENVS):
                value = clean(ws.cell(row, index + 2).value) if row else ""
                raw[component_id][env] = value
                any_value = any_value or bool(value)
        cells = {
            component_id: {
                env: cell_from(value, any_value, project_id) for env, value in envs.items()
            }
            for component_id, envs in raw.items()
        }
        projects.append(
            {
                "id": project_id,
                "name": project_id,
                "recorded": any_value,
                "cells": cells,
            }
        )
    OUT.write_text(json.dumps({"projects": projects}, indent=2) + "\n")
    recorded = [project["name"] for project in projects if project["recorded"]]
    empty = [project["name"] for project in projects if not project["recorded"]]
    print(f"wrote {OUT}")
    print("recorded", len(recorded), recorded)
    print("awaiting", len(empty), empty)


if __name__ == "__main__":
    main()
